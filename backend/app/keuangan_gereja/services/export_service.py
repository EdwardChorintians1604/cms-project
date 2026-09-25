import io
from datetime import datetime, date
from typing import List, Any, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors


class KeuanganExportService:
    """
    Layanan ekspor laporan pemasukan dan pengeluaran gereja
    menggunakan library Python openpyxl (Excel) dan ReportLab (PDF).
    """

    @staticmethod
    def _format_idr(val: float) -> str:
        return f"Rp {val:,.0f}".replace(",", ".")

    # =========================================================================
    # EXCEL EXPORT (OPENPYXL)
    # =========================================================================

    @classmethod
    def generate_excel(
        cls,
        title: str,
        items: List[Any],
        is_pemasukan: bool = True,
        church_name: str = "Seluruh Cabang Gereja"
    ) -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Laporan Keuangan"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        title_font = Font(name="Calibri", size=16, bold=True, color="78350F")
        sub_font = Font(name="Calibri", size=11, italic=True, color="475569")
        meta_font = Font(name="Calibri", size=10, bold=True, color="334155")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10, color="1E293B")
        total_font = Font(name="Calibri", size=11, bold=True, color="9A3412")

        header_fill = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
        total_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        thin_side = Side(style="thin", color="CBD5E1")
        border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        double_bottom = Border(top=thin_side, bottom=Side(style="double", color="B45309"), left=thin_side, right=thin_side)

        # Header Title Rows
        ws.merge_cells("A1:I1")
        ws["A1"] = f"GRACEPOINT CHURCH NETWORK — {title.upper()}"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A2:I2")
        ws["A2"] = f"Kriteria Cabang: {church_name} | Tanggal Ekspor: {datetime.now().strftime('%d/%m/%Y %H:%M WIB')}"
        ws["A2"].font = sub_font

        ws.append([])  # Baris 3 kosong

        # Table Column Headers
        if is_pemasukan:
            headers = ["No", "ID", "Tanggal", "Kategori", "Metode", "Penyetor / Donatur", "Status", "Keterangan", "Nominal (Rp)"]
        else:
            headers = ["No", "ID", "Tanggal", "Kategori", "Metode", "Penerima / Vendor", "Status Persetujuan", "Keterangan", "Nominal (Rp)"]

        ws.append(headers)
        header_row_idx = 4
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row_idx, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_all
        ws.row_dimensions[header_row_idx].height = 24

        # Table Data Rows
        total_nominal = 0.0
        row_idx = 5

        for idx, item in enumerate(items, start=1):
            nominal = float(getattr(item, "jumlah", 0.0))
            total_nominal += nominal

            tgl_str = str(getattr(item, "tanggal", "-"))
            kategori = str(getattr(item, "kategori", "-"))
            metode = str(getattr(item, "metode_pembayaran", "Tunai"))
            pihak = str(getattr(item, "donor_name", None) or getattr(item, "penerima", None) or "-")
            status = str(getattr(item, "status_verifikasi", None) or getattr(item, "status_persetujuan", None) or "Disetujui")
            keterangan = str(getattr(item, "keterangan", "") or "-")
            trx_id = getattr(item, "id", idx)

            row_values = [idx, trx_id, tgl_str, kategori, metode, pihak, status, keterangan, nominal]
            ws.append(row_values)

            # Apply Styles to Data Row
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.font = regular_font
                c.border = border_all
                if idx % 2 == 0:
                    c.fill = zebra_fill

                if col_idx in [1, 2, 3, 5, 7]:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 9:
                    c.alignment = Alignment(horizontal="right", vertical="center")
                    c.number_format = '"Rp "#,##0'
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

            ws.row_dimensions[row_idx].height = 20
            row_idx += 1

        # Total Row
        total_row_values = ["", "TOTAL", "", "", "", "", "", "", total_nominal]
        ws.append(total_row_values)
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.font = total_font
            c.fill = total_fill
            c.border = double_bottom
            if col_idx == 9:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = '"Rp "#,##0'
            elif col_idx == 2:
                c.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[row_idx].height = 24

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    # =========================================================================
    # PDF EXPORT (REPORTLAB)
    # =========================================================================

    @classmethod
    def generate_pdf(
        cls,
        title: str,
        items: List[Any],
        is_pemasukan: bool = True,
        church_name: str = "Seluruh Cabang Gereja"
    ) -> io.BytesIO:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )
        elements = []
        styles = getSampleStyleSheet()

        # Custom Paragraph Styles
        title_style = ParagraphStyle(
            name="DocTitle",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=16,
            textColor=colors.HexColor("#78350F"),
            spaceAfter=3,
        )
        sub_style = ParagraphStyle(
            name="DocSub",
            parent=styles["Normal"],
            fontName="Helvetica-Oblique",
            fontSize=9,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=12,
        )
        cell_style = ParagraphStyle(
            name="CellText",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8,
            textColor=colors.HexColor("#1E293B"),
            leading=10
        )
        cell_bold = ParagraphStyle(
            name="CellBold",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            textColor=colors.HexColor("#1E293B"),
            leading=10
        )

        # Header Kop Surat
        elements.append(Paragraph(f"GRACEPOINT CHURCH NETWORK — {title.upper()}", title_style))
        elements.append(Paragraph(
            f"Kriteria Cabang: <b>{church_name}</b> | Tanggal Laporan: <b>{datetime.now().strftime('%d/%m/%Y %H:%M WIB')}</b> | Soli Deo Gloria",
            sub_style
        ))

        # Headers Tabel
        if is_pemasukan:
            headers = ["No", "ID", "Tanggal", "Kategori", "Metode", "Penyetor / Donatur", "Status", "Keterangan", "Nominal (Rp)"]
        else:
            headers = ["No", "ID", "Tanggal", "Kategori", "Metode", "Penerima / Vendor", "Status", "Keterangan", "Nominal (Rp)"]

        table_data = [[Paragraph(f"<b>{h}</b>", ParagraphStyle(name="TH", fontName="Helvetica-Bold", fontSize=8, textColor=colors.white)) for h in headers]]

        total_nominal = 0.0
        for idx, item in enumerate(items, start=1):
            nominal = float(getattr(item, "jumlah", 0.0))
            total_nominal += nominal

            tgl_str = str(getattr(item, "tanggal", "-"))
            kategori = str(getattr(item, "kategori", "-"))
            metode = str(getattr(item, "metode_pembayaran", "Tunai"))
            pihak = str(getattr(item, "donor_name", None) or getattr(item, "penerima", None) or "-")
            status = str(getattr(item, "status_verifikasi", None) or getattr(item, "status_persetujuan", None) or "Disetujui")
            keterangan = str(getattr(item, "keterangan", "") or "-")
            trx_id = str(getattr(item, "id", idx))

            row = [
                Paragraph(str(idx), cell_style),
                Paragraph(trx_id, cell_style),
                Paragraph(tgl_str, cell_style),
                Paragraph(f"<b>{kategori}</b>", cell_bold),
                Paragraph(metode, cell_style),
                Paragraph(pihak, cell_style),
                Paragraph(status, cell_style),
                Paragraph(keterangan[:40] + ("..." if len(keterangan) > 40 else ""), cell_style),
                Paragraph(f"<b>{cls._format_idr(nominal)}</b>", ParagraphStyle(name="Nom", fontName="Helvetica-Bold", fontSize=8, alignment=2, textColor=colors.HexColor("#78350F"))),
            ]
            table_data.append(row)

        # Baris Total
        table_data.append([
            Paragraph("", cell_style),
            Paragraph("<b>TOTAL</b>", cell_bold),
            Paragraph("", cell_style),
            Paragraph("", cell_style),
            Paragraph("", cell_style),
            Paragraph("", cell_style),
            Paragraph("", cell_style),
            Paragraph("", cell_style),
            Paragraph(f"<b>{cls._format_idr(total_nominal)}</b>", ParagraphStyle(name="TotalNom", fontName="Helvetica-Bold", fontSize=9, alignment=2, textColor=colors.HexColor("#9A3412"))),
        ])

        # Lebar Kolom Landscape A4 (total ~ 780 pt)
        col_widths = [25, 35, 55, 110, 65, 120, 70, 180, 100]

        table = Table(table_data, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#B45309")),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F8FAFC")]),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#FEF3C7")),
            ("LINEABOVE", (0, -1), (-1, -1), 1.5, colors.HexColor("#B45309")),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
        ]))

        elements.append(table)
        elements.append(Spacer(1, 20))

        # Tanda Tangan
        sig_data = [
            [
                Paragraph("<b>Dibuat Oleh,</b><br/><br/><br/><br/><u>Bendahara Gereja / Pelayanan</u><br/>Majelis Perbendaharaan", cell_style),
                Paragraph("", cell_style),
                Paragraph("<b>Mengetahui &amp; Menyetujui,</b><br/><br/><br/><br/><u>Gembala Sidang / Ketua Sinode</u><br/>Badan Pengurus Sinode GracePoint", cell_style),
            ]
        ]
        sig_table = Table(sig_data, colWidths=[250, 260, 250])
        sig_table.setStyle(TableStyle([
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        elements.append(sig_table)

        doc.build(elements)
        buffer.seek(0)
        return buffer
