import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import json
import shutil

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    run = h.runs[0]
    if level == 1:
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(30, 64, 175) # Dark Blue
        run.bold = True
    elif level == 2:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(15, 118, 110) # Teal
        run.bold = True
    elif level == 3:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(51, 65, 85) # Slate
        run.bold = True
    return h

def create_table_spec(doc, headers, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header styling
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E40AF") # Deep blue
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(10)
            r.font.name = "Calibri"

    # Row styling
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = str(cell_value)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            for r in p.runs:
                r.font.size = Pt(9.5)
                r.font.name = "Calibri"
                if c_idx == 0:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(15, 23, 42)

    # Column widths
    col_widths = [1.5, 1.2, 1.3, 2.5] if len(headers) == 4 else [1.5, 4.5]
    for row in table.rows:
        for idx, width in enumerate(col_widths[:len(headers)]):
            row.cells[idx].width = Inches(width)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table

def add_code_block(doc, code_str):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
    cell.width = Inches(6.5)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(code_str)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def build_document():
    doc = Document()
    
    # Page setup - Margins 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Normal Style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(51, 65, 85)

    # Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("SPESIFIKASI DATABASE & PANDUAN DOCKER POSTGRESQL")
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 64, 175)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("GracePoint CMS (Content & Church Management System)")
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Divider line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(12)
    r_div = p_div.add_run("_________________________________________________________________________________")
    r_div.font.color.rgb = RGBColor(203, 213, 225)

    # 1. Pendahuluan
    add_styled_heading(doc, "1. Pendahuluan & Konsep Dasar DBMS", level=1)
    
    p = doc.add_paragraph()
    p.add_run("Dokumen ini memuat catatan spesifikasi lengkap mengenai struktur database, tabel, field, record, serta panduan deployment ")
    p.add_run("PostgreSQL menggunakan Docker").bold = True
    p.add_run(" untuk sistem GracePoint CMS. Dokumen ini dirancang sebagai panduan pengembang dalam mengimplementasikan Database Management System (DBMS) berskala industri.")

    headers_concept = ["Istilah / Elemen", "Penjelasan Konseptual", "Penerapan pada GracePoint CMS"]
    rows_concept = [
        ["Struktur Database", "Arsitektur logis dan skema keseluruhan yang mendefinisikan hubungan antar entitas.", "Skema relasional yang menghubungkan modul Jemaat, Pelayanan, Ibadah, dan Keuangan."],
        ["Tabel (Entity)", "Wadah terstruktur untuk menyimpan kumpulan objek/entitas dunia nyata yang sejenis.", "Tabel 'users', 'members', 'families', 'service_schedules', 'finances'."],
        ["Field (Column)", "Atribut spesifik yang dimiliki entitas beserta tipe data dan aturan operasinya.", "Field 'email' (VARCHAR), 'amount' (DECIMAL), 'birth_date' (DATE)."],
        ["Record (Row)", "Satu baris data utuh yang merepresentasikan satu kejadian atau objek spesifik.", "Satu baris data jemaat: ID 101, 'Hendra Kurniawan', status Aktif."],
        ["Data (Value)", "Nilai konkret yang mengisikan suatu field pada baris record tertentu.", "Nilai: 'admin@gracepoint.org', 1500000.00, true."]
    ]
    create_table_spec(doc, headers_concept, rows_concept)

    # 2. Arsitektur Relasi (ERD Overview)
    add_styled_heading(doc, "2. Ringkasan Arsitektur & Relasi Tabel (ERD)", level=1)
    p = doc.add_paragraph("Hubungan antar tabel dirancang secara terintegrasi dengan relasi utama sebagai berikut:")
    
    bullets = [
        ("users <-> members: ", "Relasi 1-to-1 (Opsional). Pengguna sistem yang merupakan jemaat terhubung ke profil anggota jemaatnya."),
        ("families <-> members: ", "Relasi 1-to-Many. Satu Kartu Keluarga (family) dapat memiliki banyak Anggota Jemaat (members)."),
        ("members <-> service_schedules: ", "Relasi 1-to-Many / Foreign Key. Anggota jemaat bertugas sebagai Worship Leader, Penyanyi, Usher, dll."),
        ("ministries <-> members: ", "Relasi 1-to-Many. Komisi pelayanan dipimpin oleh seorang Ketua Komisi (Leader)."),
        ("members <-> sacraments: ", "Relasi 1-to-Many. Pencatatan sertifikat sakramen (Baptis, Sidi, Nikah) untuk setiap anggota jemaat."),
        ("users <-> finances: ", "Relasi 1-to-Many. Setiap pencatatan transaksi keuangan merekam ID pengguna/admin yang menginput data.")
    ]
    for title, desc in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r1 = bp.add_run(title)
        r1.bold = True
        r1.font.color.rgb = RGBColor(15, 118, 110)
        bp.add_run(desc)

    # 3. Spesifikasi Detail Tabel
    add_styled_heading(doc, "3. Spesifikasi Detail Tabel, Field, & Sampel Record", level=1)
    
    tables_data = [
        {
            "name": "1. Tabel users (Akun Autentikasi Sistem)",
            "desc": "Menyimpan data kredensial dan hak akses pengguna yang dapat masuk ke aplikasi GracePoint CMS.",
            "fields": [
                ["id", "INTEGER", "PK, Autoincrement, Not Null", "ID unik pengguna"],
                ["email", "VARCHAR(255)", "Unique, Index, Not Null", "Email utama untuk login"],
                ["hashed_password", "VARCHAR(255)", "Not Null", "Password terenkripsi (Bcrypt/Argon2)"],
                ["full_name", "VARCHAR(150)", "Not Null", "Nama lengkap akun"],
                ["role", "VARCHAR(50)", "Default: 'user'", "Role: superadmin, church_admin, leader, user"],
                ["is_active", "BOOLEAN", "Default: true", "Status keaktifan akun"],
                ["created_at", "TIMESTAMP", "Default: CURRENT_TIMESTAMP", "Waktu pendaftaran akun"]
            ],
            "json_sample": {
                "id": 1,
                "email": "admin.gereja@gracepoint.org",
                "hashed_password": "$2b$12$eImiTXuWVmK8...",
                "full_name": "Pdt. Yohanes Setiawan",
                "role": "church_admin",
                "is_active": True,
                "created_at": "2026-08-02T08:00:00Z"
            }
        },
        {
            "name": "2. Tabel families (Kartu Keluarga Jemaat)",
            "desc": "Mengelompokkan anggota jemaat dalam ikatan keluarga (Kartu Keluarga Gereja).",
            "fields": [
                ["id", "INTEGER", "PK, Autoincrement, Not Null", "ID unik keluarga"],
                ["family_code", "VARCHAR(50)", "Unique, Index, Not Null", "Kode KK Gereja (misal: KK-2026-001)"],
                ["family_name", "VARCHAR(150)", "Not Null", "Nama keluarga (misal: Kel. Hendra Kurniawan)"],
                ["head_of_family_id", "INTEGER", "FK -> members.id, Nullable", "ID jemaat sebagai Kepala Keluarga"],
                ["address", "TEXT", "Nullable", "Alamat tempat tinggal keluarga"],
                ["phone_number", "VARCHAR(20)", "Nullable", "Nomor telepon rumah/keluarga"],
                ["created_at", "TIMESTAMP", "Default: CURRENT_TIMESTAMP", "Waktu pendataan keluarga"]
            ],
            "json_sample": {
                "id": 10,
                "family_code": "KK-2026-042",
                "family_name": "Keluarga Bpk. Hendra Kurniawan",
                "head_of_family_id": 101,
                "address": "Jl. Mawar No. 12, Jakarta Selatan",
                "phone_number": "021-5551234",
                "created_at": "2026-01-15T10:30:00Z"
            }
        },
        {
            "name": "3. Tabel members (Data Profil Jemaat)",
            "desc": "Menyimpan data pribadi, demografis, dan status kerohanian anggota jemaat secara rinci.",
            "fields": [
                ["id", "INTEGER", "PK, Autoincrement, Not Null", "ID unik jemaat"],
                ["user_id", "INTEGER", "FK -> users.id, Unique, Nullable", "Relasi akun login (jika ada)"],
                ["family_id", "INTEGER", "FK -> families.id, Nullable", "Relasi ke Kartu Keluarga"],
                ["member_code", "VARCHAR(50)", "Unique, Index, Not Null", "Nomor Induk Jemaat (NIJ)"],
                ["full_name", "VARCHAR(150)", "Not Null", "Nama lengkap jemaat"],
                ["gender", "VARCHAR(10)", "Check: MALE/FEMALE", "Jenis kelamin"],
                ["birth_place", "VARCHAR(100)", "Nullable", "Kota tempat lahir"],
                ["birth_date", "DATE", "Nullable", "Tanggal lahir"],
                ["phone_number", "VARCHAR(20)", "Nullable", "Nomor HP/WhatsApp"],
                ["marital_status", "VARCHAR(20)", "Default: 'SINGLE'", "SINGLE, MARRIED, DIVORCED, WIDOWED"],
                ["baptism_status", "BOOLEAN", "Default: false", "Apakah sudah dibaptis air"],
                ["confirmation_status", "BOOLEAN", "Default: false", "Apakah sudah sidi"],
                ["status", "VARCHAR(20)", "Default: 'ACTIVE'", "ACTIVE, INACTIVE, MOVED, DECEASED"],
                ["created_at", "TIMESTAMP", "Default: CURRENT_TIMESTAMP", "Waktu terdaftar"]
            ],
            "json_sample": {
                "id": 101,
                "user_id": 5,
                "family_id": 10,
                "member_code": "JMT-2026-0105",
                "full_name": "Hendra Kurniawan",
                "gender": "MALE",
                "birth_place": "Bandung",
                "birth_date": "1985-05-20",
                "phone_number": "081298765432",
                "marital_status": "MARRIED",
                "baptism_status": True,
                "confirmation_status": True,
                "status": "ACTIVE",
                "created_at": "2026-01-15T10:30:00Z"
            }
        },
        {
            "name": "4. Tabel ministries (Komisi / Divisi Pelayanan)",
            "desc": "Mendata struktur komisi pelayanan yang ada di gereja.",
            "fields": [
                ["id", "INTEGER", "PK, Autoincrement, Not Null", "ID unik komisi"],
                ["name", "VARCHAR(100)", "Not Null", "Nama komisi (Musik, Usher, Media)"],
                ["description", "TEXT", "Nullable", "Deskripsi dan fokus pelayanan"],
                ["leader_id", "INTEGER", "FK -> members.id, Nullable", "Ketua komisi pelayanan"],
                ["created_at", "TIMESTAMP", "Default: CURRENT_TIMESTAMP", "Waktu komisi dibuat"]
            ],
            "json_sample": {
                "id": 3,
                "name": "Worship Team & Musik",
                "description": "Pelayanan puji-pujian dan pengiring musik ibadah raya",
                "leader_id": 101,
                "created_at": "2026-01-10T09:00:00Z"
            }
        },
        {
            "name": "5. Tabel service_schedules (Jadwal Pelayanan Ibadah)",
            "desc": "Mendukung halaman InformasiPelayanan.vue untuk mengelola jadwal ibadah & daftar pelayan.",
            "fields": [
                ["id", "INTEGER", "PK, Autoincrement, Not Null", "ID unik jadwal"],
                ["service_date", "DATE", "Not Null", "Tanggal pelaksanaan ibadah"],
                ["service_time", "TIME", "Not Null", "Waktu/jam ibadah"],
                ["service_name", "VARCHAR(100)", "Not Null", "Nama ibadah (Ibadah Raya 1, Pemuda)"],
                ["speaker_name", "VARCHAR(150)", "Not Null", "Nama Pengkhotbah"],
                ["worship_leader_id", "INTEGER", "FK -> members.id, Nullable", "ID Worship Leader"],
                ["singers", "TEXT / JSON", "Nullable", "Daftar nama/ID penyanyi (Singer)"],
                ["musicians", "TEXT / JSON", "Nullable", "Daftar nama/ID pemain musik"],
                ["ushers", "TEXT / JSON", "Nullable", "Daftar nama/ID penyambut jemaat"],
                ["multimedia_team", "TEXT / JSON", "Nullable", "Daftar tim sound & media"],
                ["created_at", "TIMESTAMP", "Default: CURRENT_TIMESTAMP", "Waktu penginputan jadwal"]
            ],
            "json_sample": {
                "id": 50,
                "service_date": "2026-08-09",
                "service_time": "09:00:00",
                "service_name": "Ibadah Raya Mingguan 1",
                "speaker_name": "Pdt. Yohanes Setiawan",
                "worship_leader_id": 101,
                "singers": ["Maria Vania", "David Lee"],
                "musicians": ["Keyboard: Daniel", "Guitar: Kevin", "Drums: Michael"],
                "ushers": ["Sarah", "Budi", "Lani"],
                "multimedia_team": ["Operator EasyWorship: Alex", "Stream: Rizky"],
                "created_at": "2026-08-01T14:20:00Z"
            }
        },
        {
            "name": "6. Tabel church_events (Hari Besar & Event Gereja)",
            "desc": "Mendukung halaman HariBesarGereja.vue untuk mengelola kalender event dan peringatan hari besar.",
            "fields": [
                ["id", "INTEGER", "PK, Autoincrement, Not Null", "ID unik event"],
                ["title", "VARCHAR(200)", "Not Null", "Judul event / hari besar"],
                ["event_type", "VARCHAR(50)", "Not Null", "HARI_BESAR, RETREAT, SEMINAR, BAPTISAN"],
                ["start_datetime", "DATETIME", "Not Null", "Waktu & tanggal mulai"],
                ["end_datetime", "DATETIME", "Not Null", "Waktu & tanggal selesai"],
                ["location", "VARCHAR(200)", "Not Null", "Lokasi pelaksanaan acara"],
                ["description", "TEXT", "Nullable", "Rincian dan agenda event"],
                ["organizer_id", "INTEGER", "FK -> members.id, Nullable", "Ketua panitia / PIC event"],
                ["created_at", "TIMESTAMP", "Default: CURRENT_TIMESTAMP", "Waktu pendaftaran event"]
            ],
            "json_sample": {
                "id": 12,
                "title": "Ibadah & Perayaan Jumat Agung 2026",
                "event_type": "HARI_BESAR",
                "start_datetime": "2026-04-03T09:00:00Z",
                "end_datetime": "2026-04-03T12:00:00Z",
                "location": "Main Sanctuary GracePoint",
                "description": "Ibadah Peringatan Kematian Tuhan Yesus Kristus dan Perjamuan Kudus",
                "organizer_id": 101,
                "created_at": "2026-02-01T08:00:00Z"
            }
        },
        {
            "name": "7. Tabel sacraments (Catatan Sakramen)",
            "desc": "Mencatat riwayat penerimaan Sakramen Baptis, Sidi, dan Pernikahan jemaat.",
            "fields": [
                ["id", "INTEGER", "PK, Autoincrement, Not Null", "ID unik sakramen"],
                ["member_id", "INTEGER", "FK -> members.id, Not Null", "Jemaat penerima sakramen"],
                ["sacrament_type", "VARCHAR(50)", "Not Null", "BAPTISM_WATER, CONFIRMATION_SIDI, MARRIAGE"],
                ["event_date", "DATE", "Not Null", "Tanggal pelaksanaan sakramen"],
                ["officiated_by", "VARCHAR(150)", "Not Null", "Pendeta yang melayani"],
                ["certificate_number", "VARCHAR(100)", "Unique, Not Null", "Nomor sertifikat sakramen"],
                ["notes", "TEXT", "Nullable", "Catatan tambahan"],
                ["created_at", "TIMESTAMP", "Default: CURRENT_TIMESTAMP", "Waktu entri sakramen"]
            ],
            "json_sample": {
                "id": 8,
                "member_id": 101,
                "sacrament_type": "CONFIRMATION_SIDI",
                "event_date": "2020-11-15",
                "officiated_by": "Pdt. Yohanes Setiawan",
                "certificate_number": "SIDI/GP/2020/008",
                "notes": "Telah menyelesaikan bimbingan Sidi angkatan 2020",
                "created_at": "2020-11-15T11:00:00Z"
            }
        },
        {
            "name": "8. Tabel finances (Keuangan Gereja)",
            "desc": "Mencatat arus kas pemasukan persembahan/persepuluhan dan pengeluaran operasional.",
            "fields": [
                ["id", "INTEGER", "PK, Autoincrement, Not Null", "ID unik transaksi"],
                ["transaction_code", "VARCHAR(50)", "Unique, Index, Not Null", "Kode transaksi (TRX-20260802-005)"],
                ["transaction_date", "DATE", "Not Null", "Tanggal transaksi"],
                ["type", "VARCHAR(20)", "Check: INCOME / EXPENSE", "Jenis transaksi"],
                ["category", "VARCHAR(100)", "Not Null", "Persembahan, Persepuluhan, Operasional"],
                ["amount", "DECIMAL(15,2)", "Not Null", "Nominal transaksi"],
                ["description", "TEXT", "Nullable", "Keterangan rincian transaksi"],
                ["recorded_by_id", "INTEGER", "FK -> users.id, Not Null", "Admin yang mencatat"],
                ["created_at", "TIMESTAMP", "Default: CURRENT_TIMESTAMP", "Waktu input transaksi"]
            ],
            "json_sample": {
                "id": 204,
                "transaction_code": "TRX-20260802-005",
                "transaction_date": "2026-08-02",
                "type": "INCOME",
                "category": "Persembahan Ibadah Raya",
                "amount": 4750000.00,
                "description": "Persembahan Kantong Ibadah Raya Minggu 2 Agustus 2026 Sesi 1 & 2",
                "recorded_by_id": 1,
                "created_at": "2026-08-02T13:00:00Z"
            }
        },
        {
            "name": "9. Tabel announcements (Pengumuman & Warta Gereja)",
            "desc": "Menyimpan data berita dan warta jemaat yang ditayangkan di portal informasi.",
            "fields": [
                ["id", "INTEGER", "PK, Autoincrement, Not Null", "ID unik pengumuman"],
                ["title", "VARCHAR(200)", "Not Null", "Judul pengumuman"],
                ["content", "TEXT", "Not Null", "Isi pengumuman"],
                ["target_audience", "VARCHAR(50)", "Default: 'ALL'", "ALL, YOUTH, PARENTS, VOLUNTEERS"],
                ["publish_date", "DATE", "Not Null", "Tanggal mulai tayang"],
                ["expiry_date", "DATE", "Nullable", "Tanggal akhir tayang"],
                ["is_active", "BOOLEAN", "Default: true", "Status aktif tayang"],
                ["author_id", "INTEGER", "FK -> users.id, Nullable", "Pembuat warta"],
                ["created_at", "TIMESTAMP", "Default: CURRENT_TIMESTAMP", "Waktu buat warta"]
            ],
            "json_sample": {
                "id": 15,
                "title": "Pembukaan Pendaftaran Kelas Bimbingan Baptisan Air Sesi Agustus 2026",
                "content": "Diberitahukan kepada seluruh jemaat yang rindu menerima Sakramen Baptisan Air...",
                "target_audience": "ALL",
                "publish_date": "2026-08-01",
                "expiry_date": "2026-08-30",
                "is_active": True,
                "author_id": 1,
                "created_at": "2026-08-01T09:00:00Z"
            }
        }
    ]

    for tdata in tables_data:
        add_styled_heading(doc, tdata["name"], level=2)
        doc.add_paragraph(tdata["desc"])
        
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        p_lbl = doc.add_paragraph()
        r_lbl = p_lbl.add_run("Spesifikasi Field / Kolom:")
        r_lbl.bold = True
        r_lbl.font.color.rgb = RGBColor(30, 64, 175)
        
        headers_field = ["Nama Field", "Tipe Data", "Aturan / Constraint", "Deskripsi Properti"]
        create_table_spec(doc, headers_field, tdata["fields"])

        p_json = doc.add_paragraph()
        r_json = p_json.add_run("Sampel Record Data (Format JSON):")
        r_json.bold = True
        r_json.font.color.rgb = RGBColor(15, 118, 110)
        
        json_str = json.dumps(tdata["json_sample"], indent=2, ensure_ascii=False)
        add_code_block(doc, json_str)

    # 4. Panduan Docker & PostgreSQL
    add_styled_heading(doc, "4. Panduan Deployment Docker & PostgreSQL", level=1)
    p_doc = doc.add_paragraph("Berikut adalah langkah-langkah menjalankan PostgreSQL 16 & pgAdmin 4 menggunakan Docker Compose pada server/cloud:")
    
    add_styled_heading(doc, "Konfigurasi docker-compose.yml", level=2)
    docker_yml = """version: '3.8'

services:
  postgres:
    image: postgres:16-alpine
    container_name: gracepoint_postgres
    restart: always
    environment:
      POSTGRES_USER: gracepoint_user
      POSTGRES_PASSWORD: gracepoint_password
      POSTGRES_DB: gracepoint_db
    ports:
      - "5432:5432"
    volumes:
      - postgres_data:/var/lib/postgresql/data

  pgadmin:
    image: dpage/pgadmin4:latest
    container_name: gracepoint_pgadmin
    restart: always
    environment:
      PGADMIN_DEFAULT_EMAIL: admin@gracepoint.org
      PGADMIN_DEFAULT_PASSWORD: adminpassword
    ports:
      - "5050:80"

volumes:
  postgres_data:
"""
    add_code_block(doc, docker_yml)

    add_styled_heading(doc, "Konfigurasi .env pada Backend FastAPI", level=2)
    p_env = doc.add_paragraph("Ubah file backend/.env untuk mengarahkan koneksi SQLAlchemy ke PostgreSQL container:")
    env_str = """DATABASE_URL=postgresql://gracepoint_user:gracepoint_password@localhost:5432/gracepoint_db"""
    add_code_block(doc, env_str)

    add_styled_heading(doc, "Perintah Eksekusi Docker & Migrasi Database", level=2)
    cmd_str = """# 1. Jalankan PostgreSQL & pgAdmin di Docker
docker compose up -d

# 2. Pastikan kontainer berjalan
docker ps

# 3. Jalankan Migrasi Skema Alembic di folder backend
cd backend
alembic upgrade head
"""
    add_code_block(doc, cmd_str)

    # Save document
    output_path = r"d:\cms-project\documentation\Rancangan_Database_GracePoint_CMS.docx"
    doc.save(output_path)

    # Make copy in root
    shutil.copyfile(output_path, r"d:\cms-project\Rancangan_Database_GracePoint_CMS.docx")
    print("Updated docx successfully in documentation & root!")

if __name__ == "__main__":
    build_document()
