import os
import json
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

MD_CONTENT = """# PROSEDUR STANDAR PENGEMBANGAN BACK-END, API, DAN INFRASTRUKTUR
## GracePoint CMS (Content & Church Management System)
### Arsitektur Cloud-Native: Docker Hosting + Neon Serverless PostgreSQL

---

## 📋 DAFTAR ISI
1. [Bab 1: Perancangan DBMS & Arsitektur Data (Neon PostgreSQL)](#bab-1-perancangan-dbms--arsitektur-data-neon-postgresql)
2. [Bab 2: Containerization Web App & Deployment Docker](#bab-2-containerization-web-app--deployment-docker)
3. [Bab 3: Pengembangan Back-End & RESTful API (FastAPI & Python)](#bab-3-pengembangan-back-end--restful-api-fastapi--python)
4. [Bab 4: Prosedur Backup & Recovery Database](#bab-4-prosedur-backup--recovery-database)
5. [Bab 5: Manajemen Logging & Monitoring Sistem](#bab-5-manajemen-logging--monitoring-sistem)
6. [Bab 6: Manajemen Media & Storage Berkas](#bab-6-manajemen-media--storage-berkas)
7. [Bab 7: Strategi Pengujian (Testing & QA)](#bab-7-strategi-pengujian-testing--qa)
8. [Bab 8: Otomasi CI/CD & GitHub Actions (`.github`)](#bab-8-otomasi-cicd--github-actions-github)
9. [Bab 9: Checklist Operasional Pengembang](#bab-9-checklist-operasional-pengembang)

---

## Bab 1: Perancangan DBMS & Arsitektur Data (Neon PostgreSQL)

### 1.1 Keunggulan Arsitektur Neon Serverless PostgreSQL
Penggunaan **Neon PostgreSQL** (neon.tech) sebagai tempat backend database terhubung dengan web application yang di-host di Docker memberikan keuntungan besar:
- **Zero Infrastructure Overhead**: Tidak perlu mengelola container database lokal, penyimpanan disk, atau rotasi memori RAM server.
- **Enkripsi SSL/TLS Otomatis**: Koneksi terlindungi enkripsi tingkat tinggi dengan parameter `sslmode=require`.
- **Database Branching**: Kemampuan membuat skenario branch database (*Development*, *Staging*, *Production*) secara instan seperti pada Git.
- **Skalabilitas Otomatis**: Menangani lonjakan beban transaksi tanpa khawatir server down.

### 1.2 Format Connection String Neon pada Backend (`backend/.env`)
```env
# Koneksi Neon Serverless PostgreSQL
DATABASE_URL="postgresql://username:password@ep-cool-name-123456.ap-southeast-1.aws.neon.tech/gracepoint_db?sslmode=require"
```

### 1.3 Struktur Tabel Utama
Sistem memiliki 9 tabel terintegrasi:
1. `users`: Autentikasi, email, hashed password (Bcrypt), role (`superadmin`, `church_admin`, `leader`, `user`).
2. `families`: Pendataan Kartu Keluarga Gereja (`family_code`, `family_name`, `head_of_family_id`).
3. `members`: Profil jemaat, tempat/tgl lahir, status baptis/sidi, nomor induk jemaat (NIJ).
4. `ministries`: Struktur komisi & divisi pelayanan (Worship, Media, Sound, Usher, Sekolah Minggu).
5. `service_schedules`: Penjadwalan ibadah mingguan & penugasan pelayan (`InformasiPelayanan.vue`).
6. `church_events`: Event & hari besar gereja (`HariBesarGereja.vue`).
7. `sacraments`: Riwayat penerimaan sakramen (Baptis Air, Sidi, Pernikahan, Penyerahan Anak).
8. `finances`: Transaksi kas persembahan, persepuluhan, dan pengeluaran operasional.
9. `announcements`: Warta jemaat dan pengumuman sistem.

### 1.4 Manajemen Migrasi Skema ke Neon (Alembic)
```bash
cd backend
# 1. Pastikan DATABASE_URL di .env mengarah ke Neon
# 2. Jalankan migrasi skema langsung ke cloud Neon
alembic upgrade head
```

---

## Bab 2: Containerization Web App & Deployment Docker

Dengan memindahkan PostgreSQL ke **Neon Cloud**, container Docker difokuskan murni untuk meng-host **Web Application (FastAPI Backend + Vue.js Frontend)**.

### 2.1 Konfigurasi `docker-compose.yml` untuk Web Hosting
```yaml
version: '3.8'

services:
  backend:
    build:
      context: ./backend
      dockerfile: Dockerfile
    container_name: gracepoint_api
    restart: always
    ports:
      - "8000:8000"
    environment:
      - DATABASE_URL=${DATABASE_URL}
      - SECRET_KEY=${SECRET_KEY}
    volumes:
      - ./storage:/app/storage
      - ./logs:/app/logs

  frontend:
    build:
      context: ./frontend/cms-project
      dockerfile: Dockerfile
    container_name: gracepoint_web
    restart: always
    ports:
      - "80:80"
    depends_on:
      - backend
```

---

## Bab 3: Pengembangan Back-End & RESTful API (FastAPI & Python)

### 3.1 Arsitektur Folder Backend (`backend/app/`)
```
backend/app/
├── api/          # Route Handler (Endpoints REST API)
├── core/         # Config, Database Engine, Security (JWT/Bcrypt)
├── db/           # Session management & Base models
├── models/       # Model SQLAlchemy (Database entities)
├── schemas/      # Schema Pydantic (Request & Response Validation)
├── services/     # Business logic layer
└── main.py       # Entrypoint FastAPI Application
```

---

## Bab 4: Prosedur Backup & Recovery Database

Walaupun Neon memiliki fitur *Point-in-Time Restore (PITR)* bawaan, salinan backup lokal tetap disimpan di `backups/`.

### 4.1 Backup Remote dari Neon ke File Lokal (`.sql`)
```bash
pg_dump "postgresql://username:password@ep-cool-name-123456.ap-southeast-1.aws.neon.tech/gracepoint_db?sslmode=require" > backups/neon_backup_$(date +%Y%m%d_%H%M%S).sql
```

---

## Bab 5 s/d 9: Operasional & CI/CD
- **Logs**: Dicatat di `logs/app.log` dan `logs/error.log`.
- **Storage**: Media di-mount pada `storage/uploads/`.
- **Testing**: Pengujian otomatis menggunakan `pytest` di folder `backend/tests/`.
- **CI/CD**: GitHub Actions workflow pada `.github/workflows/ci.yml`.

---
*Dokumen ini diperbarui untuk mendukung arsitektur Cloud Hybrid: Docker Web Hosting + Neon Serverless PostgreSQL.*
"""

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    run = h.runs[0]
    if level == 1:
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(30, 64, 175) # Dark Blue
        run.bold = True
    elif level == 2:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(15, 118, 110) # Teal
        run.bold = True
    elif level == 3:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(51, 65, 85) # Slate
        run.bold = True
    return h

def create_table_spec(doc, headers, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header styling
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E40AF") # Deep blue
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(10)
            r.font.name = "Calibri"

    # Row styling
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = str(cell_value)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            for r in p.runs:
                r.font.size = Pt(9.5)
                r.font.name = "Calibri"
                if c_idx == 0:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table

def add_code_block(doc, code_str):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
    cell.width = Inches(6.5)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(code_str)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def generate_docx_file():
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(51, 65, 85)

    # Title
    title_p = doc.add_paragraph()
    run_title = title_p.add_run("PROSEDUR DOCKER HOSTING & NEON POSTGRESQL")
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 64, 175)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("GracePoint CMS (Content & Church Management System)")
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    p_div = doc.add_paragraph()
    r_div = p_div.add_run("_________________________________________________________________________________")
    r_div.font.color.rgb = RGBColor(203, 213, 225)

    # Bab 1
    add_styled_heading(doc, "Bab 1: Integrasi Neon Serverless PostgreSQL", level=1)
    doc.add_paragraph("Kombinasi Docker (Web App) + Neon (Cloud PostgreSQL) adalah arsitektur modern yang memisahkan beban server aplikasi dengan beban database.")
    
    headers_neon = ["Aspek Infrastruktur", "Penyedia / Platform", "Keunggulan Utama"]
    rows_neon = [
        ["Web Application", "Docker (VPS / Cloud Server)", "Ringan, terisolasi, ramah memori"],
        ["Database PostgreSQL", "Neon Cloud (neon.tech)", "Serverless, SSL Enkripsi, Zero Maintenance"],
        ["Migrasi Database", "Alembic (Python Backend)", "Skema otomatis langsung ke Cloud"],
        ["Database Backup", "Neon PITR & Remote pg_dump", "Cadangan otomatis & ekspor berkas lokal"]
    ]
    create_table_spec(doc, headers_neon, rows_neon)

    add_styled_heading(doc, "Format Connection String Neon (.env)", level=2)
    neon_env = """DATABASE_URL=postgresql://user:pass@ep-xyz.ap-southeast-1.aws.neon.tech/gracepoint_db?sslmode=require"""
    add_code_block(doc, neon_env)

    # Bab 2
    add_styled_heading(doc, "Bab 2: Docker Hosting untuk Web Application", level=1)
    doc.add_paragraph("Dengan Neon sebagai backend database, Docker Compose pada server web diatur khusus untuk aplikasi FastAPI & Vue.js:")
    docker_snippet = """docker compose up -d --build\ndocker ps"""
    add_code_block(doc, docker_snippet)

    # Save
    doc_path = r"d:\cms-project\documentation\PROSEDUR_BACKEND_DAN_API_GRACEPOINT.docx"
    doc.save(doc_path)
    shutil.copyfile(doc_path, r"d:\cms-project\PROSEDUR_BACKEND_DAN_API_GRACEPOINT.docx")

    # Also update Rancangan_Database_GracePoint_CMS.docx
    shutil.copyfile(doc_path, r"d:\cms-project\documentation\Rancangan_Database_GracePoint_CMS.docx")
    shutil.copyfile(doc_path, r"d:\cms-project\Rancangan_Database_GracePoint_CMS.docx")

def build_all():
    md_path = r"d:\cms-project\documentation\PROSEDUR_BACKEND_DAN_API_GRACEPOINT.md"
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(MD_CONTENT)
    shutil.copyfile(md_path, r"d:\cms-project\PROSEDUR_BACKEND_DAN_API_GRACEPOINT.md")

    generate_docx_file()
    print("Successfully updated docs for Neon PostgreSQL + Docker Hybrid!")

if __name__ == "__main__":
    build_all()
